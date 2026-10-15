"""Fast end-to-end MVP acceptance on an additive, isolated _test database."""

import asyncio
import io
import json
import math
import os
import secrets
import shutil
import struct
import wave
from pathlib import Path
from uuid import uuid4

import psycopg
import pytest
from cryptography.fernet import Fernet
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy.engine import make_url
from sqlalchemy.ext.asyncio import create_async_engine

from mts_backend.auth.passwords import hash_password
from mts_backend.core.config import Settings
from mts_backend.main import create_app
from mts_backend.processing.jobs import process_one

pytestmark = pytest.mark.integration
ROOT = Path(__file__).resolve().parents[3]


@pytest.fixture
def mvp(tmp_path):
    url = os.environ.get("MTS_TEST_DATABASE_URL")
    if not url:
        pytest.skip("Set MTS_TEST_DATABASE_URL to a dedicated PostgreSQL _test database.")
    if not os.environ.get("MTS_DATA_ENCRYPTION_KEY"):
        pytest.skip("Set an external MTS_DATA_ENCRYPTION_KEY for encrypted MVP checks.")
    parsed = make_url(url)
    assert (parsed.database or "").endswith("_test"), "Refusing normal database."
    sync = parsed.set(drivername="postgresql").render_as_string(hide_password=False)
    password = secrets.token_hex(10)
    org = str(uuid4())
    users = [(str(uuid4()), "mvp-" + uuid4().hex) for _ in range(2)]
    with psycopg.connect(sync) as db:
        db.execute(
            "INSERT INTO organizations(id,name,status,created_at) VALUES(%s,%s,'active',now())",
            (org, "MVP acceptance"),
        )
        for identifier, login in users:
            db.execute(
                "INSERT INTO users(id,organization_id,login,display_name,password_hash,"
                "is_active,created_at) VALUES(%s,%s,%s,%s,%s,true,now())",
                (identifier, org, login, "Acceptance user", hash_password(password)),
            )
    settings = Settings(
        _env_file=None,
        database_url=url,
        environment="test",
        cors_origins=["http://testserver"],
        cors_allow_credentials=True,
        auth_login_max_attempts=50,
        fixture_path=str(ROOT / "fixtures/dialogues_v1.json"),
        audio_storage_path=str(tmp_path / "private-audio"),
        audio_encryption_key=Fernet.generate_key().decode(),
    )
    return create_app(settings), users, password, sync, url


def signed_in(client, login, password):
    response = client.post("/api/v1/auth/login", json={"login": login, "password": password})
    assert response.status_code == 200, response.text
    return {"Origin": "http://testserver", "X-CSRF-Token": response.json()["csrf_token"]}


def changed(client, method, path, headers, body, key=None, status=200):
    headers = {**headers, **({"Idempotency-Key": key} if key else {})}
    response = client.request(method.upper(), path, json=body, headers=headers)
    assert response.status_code == status, response.text
    return response.json()


def test_demo_result_review_task_exports_and_delete(mvp):
    app, users, password, sync, _ = mvp
    with TestClient(app) as client, TestClient(create_app(app.state.settings)) as other:
        headers = signed_in(client, users[0][1], password)
        signed_in(other, users[1][1], password)
        card = changed(
            client,
            "post",
            "/api/v1/clients",
            headers,
            {"display_name": "Клиент сценария", "notes": "Private card note"},
            status=201,
        )
        assert client.get("/api/v1/clients?query=сценария").json()["total"] == 1
        create_headers = {**headers, "Idempotency-Key": uuid4().hex}
        body = {"scenario_id": "syn_01", "client_id": card["id"]}
        first = client.post("/api/v1/demo/calls", headers=create_headers, json=body)
        assert first.status_code == 202, first.text
        replay = client.post("/api/v1/demo/calls", headers=create_headers, json=body)
        assert replay.json() == first.json()
        call_id = first.json()["call_id"]
        assert other.get(f"/api/v1/calls/{call_id}/result").status_code == 404
        assert other.get(f"/api/v1/calls/{call_id}/transcript").status_code == 404
        assert client.get(f"/api/v1/clients/{card['id']}/calls").json()["total"] == 1
        result = client.get(f"/api/v1/calls/{call_id}/result")
        assert result.status_code == 200, result.text
        result = result.json()
        assert result["capabilities"]["llm"] == "disabled"
        assert result["provenance"]["is_demo"] is True
        agreement = result["agreements"][0]
        assert agreement["evidence"]
        assert client.get("/api/v1/tasks").json()["total"] == 0
        key = uuid4().hex
        acceptance = {"expected_version": agreement["version"]}
        accepted = changed(
            client, "post", f"/api/v1/agreements/{agreement['id']}/accept", headers, acceptance, key
        )
        replay = changed(
            client, "post", f"/api/v1/agreements/{agreement['id']}/accept", headers, acceptance, key
        )
        assert replay == accepted
        assert accepted["verification"] == "accepted"
        assert client.get("/api/v1/tasks").json()["total"] == 1
        assert other.get("/api/v1/tasks").json()["total"] == 0
        edited = changed(
            client,
            "patch",
            f"/api/v1/agreements/{agreement['id']}",
            headers,
            {
                "expected_version": accepted["version"],
                "action": "Проверенная задача",
                "confirm_changes": True,
            },
        )
        assert edited["verification"] == "accepted"
        history = client.get(f"/api/v1/agreements/{agreement['id']}/history")
        assert history.status_code == 200
        assert len(history.json()["items"]) >= 3
        ics = client.get("/api/v1/exports/tasks.ics?is_demo=true")
        assert ics.status_code == 200 and "BEGIN:VCALENDAR" in ics.text
        assert "DTSTART;VALUE=DATE:" in ics.text
        excel = client.get("/api/v1/exports/tasks.xlsx?is_demo=true")
        assert excel.status_code == 200
        workbook = load_workbook(io.BytesIO(excel.content))
        assert workbook.active.max_row == 2
        dashboard = client.get("/api/v1/dashboard?mode=demo").json()
        assert dashboard["agreements_count"] == 1
        assert dashboard["saved_time"]["status"] == "not_measured"
        complete = changed(
            client,
            "post",
            f"/api/v1/agreements/{agreement['id']}/execution",
            headers,
            {"expected_version": edited["version"], "action": "complete"},
            uuid4().hex,
        )
        assert complete["execution"] == "done"
        assert client.get("/api/v1/dashboard?mode=demo").json()["completion_rate"] == 1.0
        with psycopg.connect(sync) as db:
            row = db.execute("SELECT title FROM calls WHERE id=%s", (call_id,)).fetchone()
            assert row[0].startswith("fernet:v1:")
            row = db.execute(
                "SELECT data FROM agreements WHERE id=%s", (agreement["id"],)
            ).fetchone()
            assert isinstance(row[0], str) and "Проверенная задача" not in row[0]
            row = db.execute(
                "SELECT token_digest,csrf_digest FROM sessions WHERE user_id=%s", (users[0][0],)
            ).fetchone()
            assert len(row[0]) == len(row[1]) == 64
        call = client.get(f"/api/v1/calls/{call_id}").json()
        deletion = changed(
            client,
            "delete",
            f"/api/v1/calls/{call_id}",
            headers,
            {"expected_version": call["version"]},
            status=202,
        )
        assert deletion["deletion_status"] == "completed"
        assert client.get(f"/api/v1/calls/{call_id}/result").status_code == 404
        assert client.get("/api/v1/tasks").json()["total"] == 0
        assert client.get("/api/v1/exports/tasks.xlsx").status_code == 200


def tone_wav():
    output = io.BytesIO()
    with wave.open(output, "wb") as wav:
        wav.setnchannels(1)
        wav.setsampwidth(2)
        wav.setframerate(16000)
        samples = [int(1200 * math.sin(2 * math.pi * 440 * i / 16000)) for i in range(8000)]
        wav.writeframes(struct.pack("<" + "h" * len(samples), *samples))
    return output.getvalue()


def test_audio_upload_prepare_manual_transcript_provider_stub(mvp):
    if not shutil.which("ffmpeg") or not shutil.which("ffprobe"):
        pytest.skip("Install ffmpeg for the upload/preparation acceptance.")
    app, users, password, _, url = mvp
    with TestClient(app) as client:
        headers = signed_in(client, users[0][1], password)
        wav = tone_wav()
        metadata = {
            "title": "Реальное аудио без модели",
            "call_at": "2026-10-14T10:00:00+03:00",
            "timezone": "Europe/Minsk",
            "processing_rights_confirmed": True,
            "notice_version": "demo-v1",
        }
        upload_headers = {**headers, "Idempotency-Key": uuid4().hex}
        response = client.post(
            "/api/v1/calls",
            headers=upload_headers,
            data={"metadata": json.dumps(metadata)},
            files={"audio": ("unsafe-name.wav", wav, "audio/wav")},
        )
        assert response.status_code == 202, response.text
        call_id = response.json()["call_id"]
        assert client.get(f"/api/v1/calls/{call_id}/audio").content == wav
        assert all(
            path.suffix == ".enc" for path in Path(app.state.settings.audio_storage_path).iterdir()
        )
        assert wav not in next(Path(app.state.settings.audio_storage_path).iterdir()).read_bytes()

        async def prepare_once():
            engine = create_async_engine(url)
            try:
                assert await process_one(engine, app.state.settings, call_id=call_id)
            finally:
                await engine.dispose()

        asyncio.run(prepare_once())
        call = client.get(f"/api/v1/calls/{call_id}").json()
        assert call["processing_status"] == "blocked"
        assert call["analysis_status"] == "not_connected"
        transcript = changed(
            client,
            "patch",
            f"/api/v1/calls/{call_id}/transcript",
            headers,
            {
                "expected_version": call["version"],
                "segments": [{"position": 0, "text": "Подготовлю предложение к пятнице."}],
            },
            status=201,
        )
        response = client.post(
            f"/api/v1/calls/{call_id}/analysis",
            headers=headers,
            json={
                "expected_version": transcript["call_version"],
                "transcript_revision_id": transcript["id"],
            },
        )
        assert response.status_code == 503
        assert response.json()["error"]["code"] == "provider_unavailable"
        segments = transcript["segments"]
        proposal = changed(
            client,
            "post",
            f"/api/v1/calls/{call_id}/agreements",
            headers,
            {
                "action": "Подготовить предложение",
                "due": {"kind": "unknown"},
                "evidence": [
                    {
                        "transcript_revision_id": transcript["id"],
                        "segment_id": segments[0]["id"],
                        "start_cp": 0,
                        "end_cp": len(segments[0]["text"]),
                        "quote": segments[0]["text"],
                    }
                ],
            },
            status=201,
        )
        accepted = changed(
            client,
            "post",
            f"/api/v1/agreements/{proposal['id']}/accept",
            headers,
            {"expected_version": proposal["version"]},
            uuid4().hex,
        )
        assert accepted["due"]["kind"] == "unknown"
        assert client.get("/api/v1/tasks").json()["total"] == 1
        call = client.get(f"/api/v1/calls/{call_id}").json()
        changed(
            client,
            "delete",
            f"/api/v1/calls/{call_id}",
            headers,
            {"expected_version": call["version"]},
            status=202,
        )
        assert list(Path(app.state.settings.audio_storage_path).iterdir()) == []
